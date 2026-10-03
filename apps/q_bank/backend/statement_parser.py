"""
Q-Bank Forensic Statement Parsing & Extraction Engine
Extracts, cleans, and normalizes bank statement records from Excel, CSV, Word, and text buffers.
"""

import io
import os
import re
from typing import Any

import pandas as pd
from docx import Document
from loguru import logger


def extract_tables_from_word(word_file_source: Any) -> pd.DataFrame:
    """
    Extracts all table cells from a Word (.docx) document source into a normalized DataFrame.
    """
    if isinstance(word_file_source, (str, os.PathLike)):
        doc = Document(word_file_source)
    else:
        doc = Document(
            io.BytesIO(
                word_file_source.getvalue()
                if hasattr(word_file_source, "getvalue")
                else word_file_source.read()
            )
        )

    all_table_data = []
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells]
            all_table_data.append(row_text)
        all_table_data.append([])

    return pd.DataFrame(all_table_data)


def normalize_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    """
    Standardizes header column names, trims whitespace, and normalizes date/amount formats.
    """
    df = df.loc[:, ~df.columns.duplicated()].copy()
    df.columns = df.columns.astype(str).str.strip()
    col_map = {}

    possible_vdate_cols = [
        c for c in df.columns if c.lower() in ("value date", "vdate", "val date", "value_date")
    ]
    possible_date_cols = [
        c
        for c in df.columns
        if c.lower() in ("date", "txn date", "transaction date", "tran date")
        and c not in possible_vdate_cols
    ]
    possible_narr_cols = [
        c
        for c in df.columns
        if "narr" in c.lower() or "description" in c.lower() or "particular" in c.lower()
    ]
    possible_debit_cols = [
        c
        for c in df.columns
        if "withdraw" in c.lower()
        or "debit" in c.lower()
        or c.lower() in ("dr", "withdrawal", "debit amount")
    ]
    possible_credit_cols = [
        c
        for c in df.columns
        if "deposit" in c.lower()
        or "credit" in c.lower()
        or c.lower() in ("cr", "deposit", "credit amount")
    ]
    possible_bal_cols = [
        c
        for c in df.columns
        if "balance" in c.lower() or c.lower() in ("bal", "balance", "closing balance")
    ]
    possible_txnid_cols = [
        c
        for c in df.columns
        if "ref" in c.lower()
        or "chq" in c.lower()
        or "cheque" in c.lower()
        or "txn id" in c.lower()
        or "utr" in c.lower()
    ]

    possible_page_cols = [
        c
        for c in df.columns
        if "page" in c.lower() or "source_page" in c.lower() or "page_num" in c.lower()
    ]

    if possible_vdate_cols:
        col_map[possible_vdate_cols[0]] = "Value Date"
    if possible_date_cols:
        col_map[possible_date_cols[0]] = "Date"
    if possible_narr_cols:
        col_map[possible_narr_cols[0]] = "Narration"
    if possible_debit_cols:
        col_map[possible_debit_cols[0]] = "Debit Amount"
    if possible_credit_cols:
        col_map[possible_credit_cols[0]] = "Credit Amount"
    if possible_bal_cols:
        col_map[possible_bal_cols[0]] = "Closing Balance"
    if possible_txnid_cols:
        col_map[possible_txnid_cols[0]] = "Transaction ID"
    if possible_page_cols:
        col_map[possible_page_cols[0]] = "Source_Page"

    df = df.rename(columns=col_map)

    if "Source_Page" not in df.columns:
        df["Source_Page"] = "Page_1"
    else:
        df["Source_Page"] = df["Source_Page"].astype(str).fillna("Page_1")

    # Standardize Date columns
    for target_date_col in ["Date", "Value Date"]:
        if target_date_col in df.columns:
            cleaned_date_strings = []
            for raw_val in df[target_date_col].astype(str).str.strip():
                val = raw_val.split(" ")[0].replace("-", "/").strip()
                if val in ["0", "", "nan", "None"]:
                    cleaned_date_strings.append("")
                else:
                    cleaned_date_strings.append(val)
            df[target_date_col] = cleaned_date_strings

    # Standardize numerical currency amounts
    for amt_col in ["Debit Amount", "Credit Amount", "Closing Balance"]:
        if amt_col in df.columns:
            clean_s = (
                df[amt_col]
                .astype(str)
                .str.replace(",", "", regex=False)
                .str.replace("₹", "", regex=False)
                .str.replace("INR", "", regex=False)
                .str.replace(" ", "", regex=False)
                .str.strip()
            )
            df[amt_col] = pd.to_numeric(clean_s, errors="coerce").fillna(0.0)
        else:
            df[amt_col] = 0.0

    return df


def extract_clean_tracking_name(narr: str) -> str:
    """
    Forensic regex rule engine to extract clean counterparty / entity names from bank narrations.
    Isolates UPI handles, POS merchants, NEFT/RTGS sender names, and removes banking noise.
    """
    if not isinstance(narr, str) or narr.strip() == "":
        return "Internal Transfer"

    narr_clean = " ".join(narr.strip().split())
    narr_upper = narr_clean.upper()

    SYSTEM_EXCLUSIONS = {
        "UPI",
        "UPIIN",
        "UPIOUT",
        "P2A",
        "P2M",
        "PAY",
        "IN",
        "INR",
        "DR",
        "CR",
        "NEFT",
        "RTGS",
        "IMPS",
        "POS",
        "ATM",
        "TXN",
        "TRF",
        "TRANSFER",
        "NFS",
        "MAT",
        "CASH",
        "CHG",
        "CHGS",
        "REV",
        "REVERSAL",
        "DEBIT",
        "CREDIT",
        "REMITTANCE",
        "SETTLEMENT",
        "OWN",
        "WDL",
        "DEP",
        "PURCHASE",
        "YESBIFC HO",
        "SELF",
        "SWITCH",
        "AT",
        "CASHNT",
        "IW",
        "NACH",
    }

    def clean_alphabetic_only(text_segment: str) -> str:
        if "@" in text_segment:
            text_segment = text_segment.split("@")[0].strip()
        text_segment = re.sub(r"[-–—]\d+", " ", text_segment)
        text_segment = re.sub(r"\d+", " ", text_segment)
        text_segment = re.sub(
            r"\b(OKICICI|OKHDFC|OKAXIS|OKSBI|PAYTM|YBL|ABL|IPL|IBL|PHONEPE|OKHDFCBANK|OKSBB|PTYES|OKBIZAXIS|PTYBL|AXISBANK|OKCICIC|HDFCBANK|SBIBANK|ICICIBANK)\b$",
            "",
            text_segment,
            flags=re.IGNORECASE,
        )
        text_segment = text_segment.replace(".", " ").replace(",", " ")
        return " ".join(text_segment.strip().split())

    # Case 1: UPI VPA Patterns (e.g. UPI/12345/merchant@bank/Person Name)
    if "@" in narr_upper:
        tokens = [t.strip() for t in narr_upper.split("/") if t.strip()]
        for token in tokens:
            if "@" in token:
                left_side_name = token.split("@")[0].strip()
                if left_side_name.isdigit() or len(re.sub(r"\d+", "", left_side_name)) <= 2:
                    vpa_idx = tokens.index(token)
                    if vpa_idx + 1 < len(tokens):
                        next_block = tokens[vpa_idx + 1]
                        clean_next = clean_alphabetic_only(next_block)
                        if (
                            clean_next
                            and clean_next not in SYSTEM_EXCLUSIONS
                            and len(clean_next) > 2
                        ):
                            return clean_next
                else:
                    clean_vpa_user = clean_alphabetic_only(left_side_name)
                    if (
                        clean_vpa_user
                        and clean_vpa_user not in SYSTEM_EXCLUSIONS
                        and len(clean_vpa_user) > 2
                    ):
                        return clean_vpa_user

    # Case 2: Slash Delimited NetBanking (e.g. NEFT/N12345/SUPPLIER CORP/HDFC000123)
    if "/" in narr_upper:
        tokens = [t.strip() for t in narr_upper.split("/") if t.strip()]
        for token in tokens:
            clean_tok = clean_alphabetic_only(token)
            if len(clean_tok) > 2 and not any(char.isdigit() for char in clean_tok):
                if clean_tok not in SYSTEM_EXCLUSIONS:
                    if not re.match(r"^[A-Z]{4}0[A-Z0-9]{6}$", clean_tok):
                        return clean_tok

    # Case 3: POS / Card swipe merchant
    if "POS" in narr_upper or narr_upper.startswith("PURCHASE"):
        working_str = (
            narr_upper.replace("POS-", "").replace("POS", "").replace("PURCHASE", "").strip()
        )
        working_str = re.sub(r"\b\d{2}:\d{2}:\d{2}\b", " ", working_str)
        working_str = re.sub(
            r"\s+(TNIN|TN|KA|MH|DL|IN|AP|TS|MUMBAI|CHENNAI|BANGALORE)(-.*)?$", "", working_str
        )
        working_str = re.sub(r"\b[A-Z0-9]*\d+[A-Z0-9]*\b", " ", working_str)

        final_merchant = " ".join(
            [w for w in working_str.split() if w not in SYSTEM_EXCLUSIONS]
        ).strip()
        if final_merchant:
            return final_merchant

    # Fallback Case: General clean words
    scrubbed = re.sub(r"[/|\-_:+=@.]", " ", narr_upper)
    scrubbed = re.sub(r"\b\d+\b", " ", scrubbed)
    scrubbed = re.sub(r"\b[A-Z0-9]*\d+[A-Z0-9]*\b", " ", scrubbed)

    words = [w for w in scrubbed.split() if w.isalpha() and len(w) > 1]
    filtered_words = [w for w in words if w not in SYSTEM_EXCLUSIONS]

    if filtered_words:
        return " ".join(filtered_words[:3]).strip()

    return "Other Account Operations"


def format_inr(number: float | int) -> str:
    """
    Formats a numeric amount into Indian Currency notation (e.g. 1,50,000.00).
    """
    try:
        n = float(number)
    except (ValueError, TypeError):
        return "0.00"

    neg = n < 0
    n = abs(n)
    s = f"{n:.2f}"
    parts = s.split(".")
    integer_part = parts[0]
    decimal = parts[1]

    if len(integer_part) <= 3:
        formatted = integer_part
    else:
        last3 = integer_part[-3:]
        rest = integer_part[:-3]
        groups = []
        while len(rest) > 2:
            groups.append(rest[-2:])
            rest = rest[:-2]
        if rest:
            groups.append(rest)
        groups.reverse()
        formatted = ",".join(groups) + "," + last3

    if neg:
        formatted = "-" + formatted
    return formatted + "." + decimal


def parse_bank_statement_dataframe(file_obj_or_path: Any, filename: str) -> pd.DataFrame:
    """
    Master file reader routing Excel (.xlsx, .xls), CSV, Word (.docx), and tabular files
    into a standardized Pandas DataFrame with clean columns.
    """
    ext = filename.split(".")[-1].lower() if "." in filename else ""
    df = pd.DataFrame()

    if ext in ["xlsx", "xls"]:
        df = pd.read_excel(file_obj_or_path, header=0)
    elif ext == "csv":
        df = pd.read_csv(file_obj_or_path, header=0)
    elif ext == "docx":
        df = extract_tables_from_word(file_obj_or_path)
    else:
        # Attempt Excel first, fallback to CSV
        try:
            df = pd.read_excel(file_obj_or_path, header=0)
        except Exception:
            try:
                if hasattr(file_obj_or_path, "seek"):
                    file_obj_or_path.seek(0)
                df = pd.read_csv(file_obj_or_path, header=0)
            except Exception as e:
                logger.error("Failed to parse statement buffer: {}", e)
                return pd.DataFrame()

    if "page_num" in df.columns:
        df = df.drop(columns=["page_num"])

    df = normalize_dataframe(df)

    if "Name" in df.columns:
        df["UPI_Name"] = df["Name"].astype(str).str.strip()
    elif "Narration" in df.columns:
        df["UPI_Name"] = df["Narration"].apply(extract_clean_tracking_name)
    else:
        df["UPI_Name"] = "Unknown Entity"

    df["Name"] = df["UPI_Name"]
    return df


def extract_statement_period(df: pd.DataFrame) -> str:
    """
    Extracts the statement period (min date - max date) from the transaction table DataFrame.
    Returns a human-readable date range string (e.g. '01 Mar 2026 - 28 Mar 2026') or empty string.
    """
    if df is None or df.empty:
        return ""

    dates = []
    for col in ["Date", "Value Date", "Txn Date", "Transaction Date", "Posting Date"]:
        if col in df.columns:
            for val in df[col]:
                raw_s = str(val).strip()
                if not raw_s or raw_s.lower() in ["0", "nan", "none", "—", "-"]:
                    continue
                try:
                    if re.match(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}", raw_s):
                        dt = pd.to_datetime(raw_s, dayfirst=False, errors="coerce")
                    else:
                        dt = pd.to_datetime(raw_s, dayfirst=True, errors="coerce")
                    if pd.notna(dt):
                        dates.append(dt)
                except Exception:
                    continue

    if not dates:
        return ""

    min_date = min(dates)
    max_date = max(dates)

    if pd.isna(min_date) or pd.isna(max_date):
        return ""

    fmt = "%d %b %Y"
    if min_date == max_date:
        return min_date.strftime(fmt)
    return f"{min_date.strftime(fmt)} - {max_date.strftime(fmt)}"

